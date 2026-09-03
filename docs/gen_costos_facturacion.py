from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "Simulador de Costos"

hf = Font(name="Arial", bold=True, color="FFFFFF", size=11)
hfill = PatternFill("solid", fgColor="1a3764")
ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
b = Border(left=Side("thin", color="D1D5DB"), right=Side("thin", color="D1D5DB"), top=Side("thin", color="D1D5DB"), bottom=Side("thin", color="D1D5DB"))
nf = Font(name="Arial", size=10)
bf = Font(name="Arial", bold=True, size=10)
gf = Font(name="Arial", bold=True, color="166534", size=10)
ca = Alignment(horizontal="center", vertical="center", wrap_text=True)
la = Alignment(vertical="center", wrap_text=True)
green_fill = PatternFill("solid", fgColor="DCFCE7")
blue_fill = PatternFill("solid", fgColor="DBEAFE")
yellow_fill = PatternFill("solid", fgColor="FEF3C7")

# TITULO
ws.merge_cells("A1:E1")
ws["A1"] = "WhatsApp Business API - Categorias de Mensajes y Costos (Peru)"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color="1a3764")
ws.merge_cells("A2:E2")
ws["A2"] = "Referencia para integracion de facturacion electronica via WhatsApp"
ws["A2"].font = Font(name="Arial", italic=True, size=10, color="6B7280")

# TABLA CATEGORIAS
cat_headers = ["Categoria", "Quien inicia?", "Ejemplo", "Costo Peru (USD)", "Necesita plantilla?"]
for c, h in enumerate(cat_headers, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = hf
    cell.fill = hfill
    cell.alignment = ha
    cell.border = b

cat_data = [
    ["Servicio", "El CLIENTE escribe primero", "Cliente: hola -> Bot responde reclamos, consultas, asesor", "GRATIS", "No"],
    ["Utilidad", "La EMPRESA envia primero", "Comprobante de pago, estado de reclamo, confirmacion de cita", "$0.0200", "Si (aprobada por Meta)"],
    ["Marketing", "La EMPRESA envia primero", "Promociones, ofertas, newsletters, campanas", "$0.0703", "Si (aprobada por Meta)"],
    ["Autenticacion", "La EMPRESA envia primero", "Codigo OTP, verificacion de cuenta, 2FA", "$0.0200", "Si (aprobada por Meta)"],
]

for r, rd in enumerate(cat_data, 5):
    for c, val in enumerate(rd, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = b
        cell.alignment = la if c <= 3 else ca
        cell.font = nf
        if r == 5:
            cell.fill = green_fill
            if c == 4:
                cell.font = gf
        elif r == 6:
            cell.fill = yellow_fill
            if c == 4:
                cell.font = bf

# SIMULADOR
ws.merge_cells("A11:E11")
ws["A11"] = "Simulador de Costos - Facturacion Electronica por WhatsApp"
ws["A11"].font = Font(name="Arial", bold=True, size=13, color="1a3764")

sim_headers = ["Volumen mensual", "Costo por mensaje", "Costo mensual (USD)", "Costo mensual (PEN)", "Costo anual (USD)"]
for c, h in enumerate(sim_headers, 1):
    cell = ws.cell(row=13, column=c, value=h)
    cell.font = hf
    cell.fill = hfill
    cell.alignment = ha
    cell.border = b

volumes = [100, 250, 500, 1000, 2000, 5000, 10000, 25000, 50000]
rate = 0.02
pen_rate = 3.75

for i, vol in enumerate(volumes):
    r = 14 + i
    ws.cell(row=r, column=1, value=vol).font = bf
    ws.cell(row=r, column=1).alignment = ca
    ws.cell(row=r, column=1).number_format = "#,##0"
    ws.cell(row=r, column=2, value=rate).number_format = "$#,##0.0000"
    ws.cell(row=r, column=2).alignment = ca
    ws.cell(row=r, column=2).font = nf
    ws.cell(row=r, column=3, value=f"=A{r}*B{r}")
    ws.cell(row=r, column=3).number_format = "$#,##0.00"
    ws.cell(row=r, column=3).alignment = ca
    ws.cell(row=r, column=3).font = bf
    ws.cell(row=r, column=4, value=f"=C{r}*{pen_rate}")
    ws.cell(row=r, column=4).number_format = "S/ #,##0.00"
    ws.cell(row=r, column=4).alignment = ca
    ws.cell(row=r, column=4).font = bf
    ws.cell(row=r, column=5, value=f"=C{r}*12")
    ws.cell(row=r, column=5).number_format = "$#,##0.00"
    ws.cell(row=r, column=5).alignment = ca
    ws.cell(row=r, column=5).font = nf
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = b
    if vol == 1000:
        for c in range(1, 6):
            ws.cell(row=r, column=c).fill = yellow_fill

# COMPARATIVA
ws.merge_cells("A25:E25")
ws["A25"] = "Comparativa: Su sistema actual vs Integracion de facturacion"
ws["A25"].font = Font(name="Arial", bold=True, size=13, color="1a3764")

comp_headers = ["Concepto", "Sistema actual (Bot de reclamos)", "Con facturacion electronica"]
for c, h in enumerate(comp_headers, 1):
    cell = ws.cell(row=27, column=c, value=h)
    cell.font = hf
    cell.fill = hfill
    cell.alignment = ha
    cell.border = b

comp_data = [
    ["Quien inicia?", "El cliente siempre", "La empresa (envio de comprobante)"],
    ["Categoria Meta", "Servicio", "Utilidad"],
    ["Costo por mensaje", "GRATIS", "$0.02 USD por comprobante enviado"],
    ["Necesita plantilla?", "No", "Si (debe ser aprobada por Meta)"],
    ["Ventana de 24h aplica", "Si (gratis dentro de ventana)", "No aplica (es mensaje proactivo)"],
    ["Ejemplo", 'Cliente: "quiero registrar un reclamo"', 'Empresa: "Su boleta B001-00542 por S/150.00"'],
    ["Costo mensual (1,000 msg)", "$0.00", "~$20.00 USD (~S/75.00)"],
]

for r, rd in enumerate(comp_data, 28):
    for c, val in enumerate(rd, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = b
        cell.alignment = la
        cell.font = nf
        if c == 2:
            cell.fill = green_fill
        if c == 3:
            cell.fill = blue_fill
        if val in ("GRATIS", "$0.00"):
            cell.font = gf

# NOTAS
ws.merge_cells("A37:E37")
ws["A37"] = "NOTA: Si el cliente ya escribio y esta dentro de la ventana de 24h, se puede enviar el comprobante GRATIS como respuesta (categoria Servicio)."
ws["A37"].font = Font(name="Arial", bold=True, italic=True, size=10, color="92400E")
ws["A37"].fill = PatternFill("solid", fgColor="FEF3C7")
ws["A37"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[37].height = 35

ws.merge_cells("A38:E38")
ws["A38"] = "Tipo de cambio referencial: 1 USD = S/ 3.75 PEN. Los precios de Meta pueden variar. Consultar: business.whatsapp.com/products/platform-pricing"
ws["A38"].font = Font(name="Arial", size=9, color="6B7280")

ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 35
ws.column_dimensions["C"].width = 35
ws.column_dimensions["D"].width = 22
ws.column_dimensions["E"].width = 22

# === HOJA 2: Flujo de Decision ===
ws2 = wb.create_sheet("Flujo de Decision")
ws2.merge_cells("A1:C1")
ws2["A1"] = "Cuando es gratis y cuando se cobra?"
ws2["A1"].font = Font(name="Arial", bold=True, size=14, color="1a3764")

flow_headers = ["Escenario", "Categoria", "Costo"]
for c, h in enumerate(flow_headers, 1):
    cell = ws2.cell(row=3, column=c, value=h)
    cell.font = hf
    cell.fill = hfill
    cell.alignment = ha
    cell.border = b

flows = [
    ['Cliente escribe "hola" -> Bot responde', "Servicio", "GRATIS"],
    ["Cliente consulta estado de reclamo -> Bot responde", "Servicio", "GRATIS"],
    ["Cliente pide hablar con asesor -> Asesor responde", "Servicio", "GRATIS"],
    ["Asesor responde dentro de 24h desde ultimo mensaje del cliente", "Servicio", "GRATIS"],
    ["Cliente escribe -> Bot registra reclamo -> Envia confirmacion", "Servicio", "GRATIS"],
    ["Empresa envia comprobante SIN que cliente haya escrito", "Utilidad", "$0.02 USD"],
    ["Empresa envia notificacion de estado SIN que cliente haya escrito", "Utilidad", "$0.02 USD"],
    ["Empresa envia promocion o campana", "Marketing", "$0.0703 USD"],
    ["Empresa envia codigo de verificacion OTP", "Autenticacion", "$0.02 USD"],
    ["Cliente escribio hace < 24h -> Empresa responde con comprobante", "Servicio", "GRATIS"],
]

for r, rd in enumerate(flows, 4):
    for c, val in enumerate(rd, 1):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.border = b
        cell.font = nf
        cell.alignment = la if c == 1 else ca
        if val == "GRATIS":
            cell.font = gf
            cell.fill = green_fill
        elif val == "Servicio":
            cell.fill = green_fill
        elif "USD" in str(val) and val != "GRATIS":
            cell.font = Font(name="Arial", bold=True, color="92400E", size=10)
            cell.fill = PatternFill("solid", fgColor="FEF3C7")
        elif val in ("Utilidad", "Marketing", "Autenticacion"):
            cell.fill = yellow_fill

ws2.column_dimensions["A"].width = 58
ws2.column_dimensions["B"].width = 18
ws2.column_dimensions["C"].width = 18

out = "C:/Users/user/Desktop/SaasLibroReclamaciones/docs/WhatsApp_Costos_Facturacion_Electronica.xlsx"
wb.save(out)
print("OK:", out)
